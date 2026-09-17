"""Build the management Word report + numbered-step diagram PNGs + Excalidraw files."""

from __future__ import annotations

import json
import uuid
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
PNG_DIR = ROOT / "diagrams"
EXCAL_DIR = ROOT / "excalidraw"
DOCX = ROOT / "Buyer-Category-and-Global-Stocks.docx"
LOGO = Path(__file__).resolve().parents[2] / "stock_ledger" / "assets" / "gazebo-logo.png"

ORANGE = (232, 119, 34)
NAVY = (31, 42, 55)
SLATE = (71, 85, 105)
WHITE = (255, 255, 255)
CREAM = (255, 247, 237)
GREEN = (22, 101, 52)
RED = (153, 27, 27)

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_B if bold else FONT, size)


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for w in words:
        trial = f"{cur} {w}".strip()
        if draw.textlength(trial, font=fnt) <= width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def _id() -> str:
    return uuid.uuid4().hex[:20]


def save_excalidraw(name: str, title: str, steps: list[str]) -> None:
    elements = []
    elements.append(
        {
            "id": _id(),
            "type": "text",
            "x": 40,
            "y": 24,
            "width": 900,
            "height": 40,
            "text": title,
            "originalText": title,
            "fontSize": 24,
            "fontFamily": 1,
            "textAlign": "left",
            "verticalAlign": "top",
            "strokeColor": "#1f2a37",
            "backgroundColor": "transparent",
            "fillStyle": "solid",
            "strokeWidth": 1,
            "roughness": 1,
            "opacity": 100,
            "angle": 0,
            "seed": 1,
            "version": 1,
            "versionNonce": 1,
            "isDeleted": False,
            "boundElements": None,
            "updated": 1,
            "link": None,
            "locked": False,
            "containerId": None,
            "lineHeight": 1.25,
            "autoResize": True,
        }
    )
    x = 40
    y = 90
    prev = None
    row = 0
    for i, step in enumerate(steps, start=1):
        if i > 1 and (i - 1) % 4 == 0:
            row += 1
            x = 40
            y = 90 + row * 200
            prev = None
        box_id = _id()
        text_id = _id()
        label = f"{i}. {step}"
        elements.append(
            {
                "id": box_id,
                "type": "rectangle",
                "x": x,
                "y": y,
                "width": 260,
                "height": 140,
                "strokeColor": "#e87722",
                "backgroundColor": "#fff7ed",
                "fillStyle": "solid",
                "strokeWidth": 2,
                "roughness": 1,
                "opacity": 100,
                "angle": 0,
                "seed": i,
                "version": 1,
                "versionNonce": i,
                "isDeleted": False,
                "boundElements": [{"id": text_id, "type": "text"}],
                "updated": 1,
                "link": None,
                "locked": False,
                "roundness": {"type": 3},
            }
        )
        elements.append(
            {
                "id": text_id,
                "type": "text",
                "x": x + 16,
                "y": y + 24,
                "width": 228,
                "height": 90,
                "text": label,
                "originalText": label,
                "fontSize": 18,
                "fontFamily": 1,
                "textAlign": "left",
                "verticalAlign": "top",
                "strokeColor": "#1f2a37",
                "backgroundColor": "transparent",
                "fillStyle": "solid",
                "strokeWidth": 1,
                "roughness": 1,
                "opacity": 100,
                "angle": 0,
                "seed": i + 50,
                "version": 1,
                "versionNonce": i + 50,
                "isDeleted": False,
                "boundElements": None,
                "updated": 1,
                "link": None,
                "locked": False,
                "containerId": box_id,
                "lineHeight": 1.25,
                "autoResize": True,
            }
        )
        if prev is not None:
            elements.append(
                {
                    "id": _id(),
                    "type": "arrow",
                    "x": prev + 260,
                    "y": y + 70,
                    "width": 40,
                    "height": 0,
                    "strokeColor": "#e87722",
                    "backgroundColor": "transparent",
                    "fillStyle": "solid",
                    "strokeWidth": 2,
                    "roughness": 1,
                    "opacity": 100,
                    "angle": 0,
                    "seed": i + 80,
                    "version": 1,
                    "versionNonce": i + 80,
                    "isDeleted": False,
                    "boundElements": None,
                    "updated": 1,
                    "link": None,
                    "locked": False,
                    "points": [[0, 0], [40, 0]],
                    "lastCommittedPoint": None,
                    "startBinding": None,
                    "endBinding": None,
                    "startArrowhead": None,
                    "endArrowhead": "arrow",
                    "elbowed": False,
                }
            )
        prev = x
        x += 320
    payload = {
        "type": "excalidraw",
        "version": 2,
        "source": "https://excalidraw.com",
        "elements": elements,
        "appState": {"gridSize": None, "viewBackgroundColor": "#ffffff"},
        "files": {},
    }
    EXCAL_DIR.mkdir(parents=True, exist_ok=True)
    (EXCAL_DIR / f"{name}.excalidraw").write_text(json.dumps(payload, indent=2))


def flow(name: str, title: str, steps: list[str], footnote: str = "") -> Path:
    cols = min(len(steps), 4)
    rows = (len(steps) + 3) // 4
    W = max(980, 80 + cols * 340)
    H = 150 + rows * 210 + (70 if footnote else 0)
    img = Image.new("RGB", (W, H), WHITE)
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, W, 8], fill=ORANGE)
    draw.text((40, 28), title, font=font(26, True), fill=NAVY)

    box_w, box_h = 280, 150
    for i, step in enumerate(steps):
        col = i % 4
        row = i // 4
        x = 40 + col * 340
        y = 90 + row * 210
        draw.rounded_rectangle(
            [x, y, x + box_w, y + box_h], radius=16, fill=CREAM, outline=ORANGE, width=3
        )
        draw.ellipse([x + 14, y + 14, x + 48, y + 48], fill=ORANGE)
        n = str(i + 1)
        nf = font(18, True)
        tw = draw.textlength(n, font=nf)
        draw.text((x + 31 - tw / 2, y + 20), n, font=nf, fill=WHITE)
        lines = wrap(draw, step, font(18), box_w - 28)
        ty = y + 62
        for line in lines[:4]:
            draw.text((x + 14, ty), line, font=font(18), fill=NAVY)
            ty += 24
        if col < 3 and i + 1 < len(steps) and (i + 1) % 4 != 0:
            ax = x + box_w + 8
            ay = y + box_h // 2
            draw.polygon([(ax, ay - 8), (ax + 28, ay), (ax, ay + 8)], fill=ORANGE)

    if footnote:
        for i, line in enumerate(wrap(draw, footnote, font(16), W - 80)):
            draw.text((40, H - 56 + i * 22), line, font=font(16), fill=SLATE)

    PNG_DIR.mkdir(parents=True, exist_ok=True)
    path = PNG_DIR / f"{name}.png"
    img.save(path, "PNG")
    save_excalidraw(name, title, steps)
    return path


def before_after(
    name: str,
    title: str,
    left_title: str,
    left_rows: list[str],
    right_title: str,
    right_rows: list[str],
    steps: list[str],
    rule: str,
) -> Path:
    W, H = 1600, 780
    img = Image.new("RGB", (W, H), WHITE)
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, W, 8], fill=ORANGE)
    draw.text((40, 24), title, font=font(26, True), fill=NAVY)

    def panel(x, y, w, h, heading, rows, color):
        draw.rounded_rectangle(
            [x, y, x + w, y + h], radius=16, fill=CREAM, outline=color, width=3
        )
        draw.text((x + 20, y + 16), heading, font=font(20, True), fill=color)
        ty = y + 58
        for r in rows:
            draw.text((x + 20, ty), r, font=font(18), fill=NAVY)
            ty += 32

    panel(40, 90, 620, 280, left_title, left_rows, RED)
    panel(940, 90, 620, 280, right_title, right_rows, GREEN)
    draw.polygon([(700, 220), (880, 250), (700, 280)], fill=ORANGE)
    draw.text((740, 188), "then", font=font(16, True), fill=ORANGE)

    draw.text((40, 410), "Steps", font=font(20, True), fill=NAVY)
    x = 40
    y = 450
    for i, step in enumerate(steps, start=1):
        draw.ellipse([x, y, x + 36, y + 36], fill=ORANGE)
        draw.text((x + 12, y + 6), str(i), font=font(16, True), fill=WHITE)
        draw.rounded_rectangle(
            [x + 48, y - 6, x + 360, y + 42], radius=10, fill=CREAM, outline=ORANGE, width=2
        )
        draw.text((x + 60, y + 4), step, font=font(16), fill=NAVY)
        if i < len(steps):
            draw.polygon(
                [(x + 372, y + 12), (x + 396, y + 18), (x + 372, y + 24)], fill=ORANGE
            )
        x += 390

    draw.rounded_rectangle([40, 540, 1560, 720], radius=14, fill=(254, 242, 242), outline=RED, width=2)
    draw.text((60, 560), "Rule", font=font(18, True), fill=RED)
    for i, line in enumerate(wrap(draw, rule, font(20, True), 1440)):
        draw.text((60, 600 + i * 32), line, font=font(20, True), fill=NAVY)

    PNG_DIR.mkdir(parents=True, exist_ok=True)
    path = PNG_DIR / f"{name}.png"
    img.save(path, "PNG")
    save_excalidraw(name, title, steps + [rule])
    return path


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    r = run._element.get_or_add_rPr()
    rFonts = r.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        r.append(rFonts)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "CBD5E1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_p(doc, text, size=11, bold=False, color=None, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True, color=ORANGE)
    return p


def subhead(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=NAVY)
    return p


def bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=11, color=NAVY)
        p.paragraph_format.space_after = Pt(2)


def picture(doc, path: Path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.4))


def evidence(doc):
    subhead(doc, "Evidence of outcome (after developer)")
    add_p(
        doc,
        "Leave this blank for now. After the developer finishes, record your screen and paste the file or a link in the row below.",
        size=10,
        color=SLATE,
        space_after=6,
    )
    rows = [
        ("Date recorded", ""),
        ("Recorded by", ""),
        ("Screen recording / screenshot", ""),
        ("Pass / Fail", ""),
        ("Notes", ""),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].paragraphs[0].clear()
    r = hdr[0].paragraphs[0].add_run("Field")
    set_run_font(r, size=10, bold=True, color=WHITE)
    hdr[1].paragraphs[0].clear()
    r2 = hdr[1].paragraphs[0].add_run("Fill in after the work is done")
    set_run_font(r2, size=10, bold=True, color=WHITE)
    shade(hdr[0], "E87722")
    shade(hdr[1], "E87722")
    set_cell_border(hdr[0])
    set_cell_border(hdr[1])
    for i, (label, _) in enumerate(rows, start=1):
        c0, c1 = table.rows[i].cells
        c0.paragraphs[0].clear()
        rr = c0.paragraphs[0].add_run(label)
        set_run_font(rr, size=10, bold=True, color=NAVY)
        c1.paragraphs[0].clear()
        blank = c1.paragraphs[0].add_run("________________________________")
        set_run_font(blank, size=10, color=SLATE)
        shade(c0, "FFF7ED")
        set_cell_border(c0)
        set_cell_border(c1)
    doc.add_paragraph()


def caption(doc, text):
    p = add_p(doc, text, size=9, color=SLATE, space_after=10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_diagrams() -> dict[str, Path]:
    d: dict[str, Path] = {}
    d["big"] = flow(
        "01-big-picture",
        "The whole change, in order",
        [
            "Admin opens a raw material in Product",
            "Sets Buyer category on Purchase Guide",
            "Global Stocks can filter by product category and buyer category",
            "User picks which stock columns to see",
            "Open purchase orders show on the right of the row",
            "Each night at 00:10 an email lists POs still open",
        ],
        "Planning Excel can also add one extra sheet that totals the same raw material.",
    )
    d["buyer"] = flow(
        "02-set-buyer-category",
        "Req 1 — set Buyer category",
        [
            "Open the product (raw material)",
            "Go to Purchase, then Guide",
            "New row: Buyer category",
            "Pick Key product, Other product, or Planner monitor",
            "Save. The timeline records who changed it",
        ],
    )
    d["audit"] = flow(
        "03-audit-log",
        "Req 1 — every change is recorded",
        [
            "Admin changes Buyer category",
            "System stores the old value and the new value",
            "Timeline shows who, when, and what changed",
            "This can be changed again later. Each change is a new log line",
        ],
    )
    d["filters"] = flow(
        "04-stock-filters",
        "Req 2 — Global Stocks filters",
        [
            "Choose Category (same as Products page)",
            "Choose Sub-category",
            "Choose Sub-sub category",
            "Choose Buyer category (Key / Other / Planner monitor)",
            "Keep using Location as today (Unit 2, Low Risk, and so on)",
        ],
        "Search, Item, Warning and Hide zero stay as they are.",
    )
    d["chooser"] = flow(
        "05-column-chooser",
        "Req 3 — pick which columns you see",
        [
            "Look at the stock table header",
            "Right-click the header",
            "Tick a column to show it, or untick to hide it",
            "Scroll right to see extra columns such as purchase orders",
        ],
    )
    d["peas"] = before_after(
        "06-peas-trace-rollup",
        "Req 3.1 — hide Trace number, add qty",
        "Before (Trace is visible)",
        [
            "Peas | Unit 2 | Trace A | 10 kg",
            "Peas | Unit 2 | Trace B | 8 kg",
            "Two rows. Same product. Same location.",
        ],
        "After (Trace is hidden)",
        [
            "Peas | Unit 2 | 18 kg",
            "One row. Qty 10 + 8 = 18 kg.",
            "Product code stays the key. Name stays with it.",
        ],
        ["Right-click the header", "Untick Trace number", "Qty from those lots is added"],
        "Locations never merge. Unit 2 and Low Risk stay separate even when Trace is hidden.",
    )
    d["salt"] = before_after(
        "07-salt-supplier-rollup",
        "Req 3.1 — hide Supplier, add qty",
        "Before (Supplier is visible)",
        [
            "Salt | Unit 2 | Supplier A | 5 kg",
            "Salt | Unit 2 | Supplier B | 5 kg",
            "Salt | Unit 2 | Supplier C | 5 kg",
            "Salt | Unit 2 | Supplier D | 5 kg",
        ],
        "After (Supplier is hidden)",
        [
            "Salt | Unit 2 | 20 kg",
            "One row. 5 + 5 + 5 + 5 = 20 kg.",
            "Still only that location.",
        ],
        ["Right-click the header", "Untick Supplier", "Qty from those suppliers is added"],
        "Do not mix Unit 2 stock with Low Risk stock. Warehouses stay split.",
    )
    d["noloc"] = flow(
        "08-never-merge-locations",
        "Hard rule — do not mix locations",
        [
            "Peas in Unit 2 stay on a Unit 2 row",
            "Peas in Low Risk stay on a Low Risk row",
            "Hiding Trace or Supplier only adds qty inside the same location",
            "Grand total columns follow the same split",
        ],
    )
    d["po"] = flow(
        "09-future-pos",
        "Req 4 — future purchase orders on the row",
        [
            "Find the product on Global Stocks (example: peas)",
            "Scroll right to the purchase order columns",
            "PO 1 is the soonest open order: remaining qty, delivery date, PO number",
            "PO 2, PO 3 and so on sit next to it, in date order",
        ],
        "Draft, cancelled and fully received POs do not show.",
    )
    d["slide"] = flow(
        "10-goods-in-slides-pos",
        "Req 4.2 — goods-in moves the list",
        [
            "PO 43123 ordered 90 kg peas. It sits in PO 1",
            "Warehouse receives 40 kg. Status is Partial. Qty shows 50 kg left",
            "Warehouse receives the last 50 kg. That PO is finished",
            "PO 43123 drops off the report",
            "The next future PO slides into PO 1",
        ],
    )
    d["totals"] = flow(
        "11-three-qty-columns",
        "Req 5 — three quantity columns",
        [
            "Current stock qty (already on the page)",
            "Future order stock qty = remaining qty on open POs",
            "Grand total = current stock + future order stock",
        ],
        "Example: 100 kg in stock + 50 kg still on POs = 150 kg grand total.",
    )
    d["email"] = flow(
        "12-night-email",
        "Req 6 — reminder email at 00:10",
        [
            "Admin sets the email addresses in the app",
            "Every night at 00:10 the job runs",
            "It lists POs that are Ordered or Partial",
            "Newest PO date is first (16 Sep, then 15 Sep, then 14 Sep)",
            "Email looks like the closing stock report. A file is attached",
        ],
        "Draft, received and cancelled POs are left out.",
    )
    d["excel"] = flow(
        "13-planning-excel",
        "Req 7 — consolidate raw material in Excel",
        [
            "Open a plan and click Export Excel",
            "A popup asks: Consolidate req of raw material?",
            "If you tick it, a new sheet is added",
            "Peas 2 kg at one level + 10 kg at another level become Peas 12 kg",
        ],
        "The two Excel sheets you already have stay the same.",
    )
    d["who"] = flow(
        "14-who-does-what",
        "Who does each step",
        [
            "Admin sets Buyer category on each raw material",
            "Buyer uses that label when placing orders later",
            "Warehouse does goods-in. Finished POs leave the stock report",
            "Planner can export a consolidated raw-material total",
            "Management checks this pack and the screen recordings",
        ],
    )
    d["safe"] = flow(
        "15-what-stays-the-same",
        "What we will not break",
        [
            "Purchase Guide fields that already exist still save the same way",
            "Global Stocks live update, Location, Item, Warning and Hide zero stay",
            "Default stock columns stay the first view",
            "Planning Excel sheets Plan Requirements and Working stay",
            "Closing stock email and its address list stay separate",
        ],
    )
    return d


def build_doc(d: dict[str, Path]) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Inches(1.6))

    add_p(doc, "GAZEBO CLOUD  ·  MANAGEMENT PACK", size=10, bold=True, color=ORANGE, space_after=4)
    add_p(
        doc,
        "Buyer Category, Global Stocks, and Purchase Order Follow-up",
        size=22,
        bold=True,
        color=NAVY,
        space_after=6,
    )
    add_p(
        doc,
        "Simple English pack for management. Diagrams use numbered steps. Evidence boxes are blank so you can paste a screen recording after the developer finishes.",
        size=12,
        color=SLATE,
        space_after=8,
    )
    bullets(
        doc,
        [
            "For: Management",
            "Date: 17 September 2026",
            "Status: Proposed — not built yet",
            "How to use: read each page, look at the numbered picture, leave the evidence table empty until you record the screen",
        ],
    )

    heading(doc, "1. Why we are doing this")
    add_p(
        doc,
        "Buyers and planners need a simple label on raw materials: Key product, Other product, or Planner monitor. Admin sets that label on the product. It can be changed later. Every change is kept in the product timeline.",
    )
    add_p(
        doc,
        "Global Stocks should filter like the Products page (category, sub-category, sub-sub category), plus this new buyer label. The Location filter we already have stays. People should also be able to tick which columns they see. If they hide Trace or Supplier, quantities add up — but Unit 2 and Low Risk never mix.",
    )
    add_p(
        doc,
        "Open purchase orders sit on the right of that report. When warehouse finishes goods-in for a PO, that PO leaves the report and the next one moves up. Each night at 00:10 we email the list of POs still Ordered or Partial. Planning Excel can optionally add one sheet that totals the same raw material across recipe levels.",
    )
    picture(doc, d["big"])
    caption(doc, "Picture 1. The whole change, step by step.")

    heading(doc, "2. Who does what")
    picture(doc, d["who"])
    caption(doc, "Picture 2. Admin, buyer, warehouse, planner, management.")

    heading(doc, "3. Req 1 — Buyer category")
    add_p(doc, "What it is. Three labels on a raw material.")
    bullets(
        doc,
        [
            "Key product — used a lot, important every day",
            "Other product — not a key product",
            "Planner monitor — the production planner watches it and orders when needed",
        ],
    )
    add_p(doc, "Where it lives. Product → Purchase → Guide. New row: Buyer category.")
    add_p(doc, "Who uses it. Admin sets it. Buyers use it later when placing orders.")
    add_p(doc, "It can be edited any time. The timeline records the old value, the new value, who, and when.")
    picture(doc, d["buyer"])
    caption(doc, "Picture 3. How admin sets Buyer category.")
    picture(doc, d["audit"])
    caption(doc, "Picture 4. How the audit log stores each change.")
    add_p(doc, "What “done” looks like. The dropdown saves. The timeline shows the change. Old Purchase Guide fields still work.")
    evidence(doc)

    heading(doc, "4. Req 2 — Filters on Global Stocks")
    add_p(doc, "Add the same category menus as the Products page, plus Buyer category. Keep Location.")
    bullets(
        doc,
        [
            "Category → Sub-category → Sub-sub category",
            "Buyer category (the three labels from Req 1)",
            "Location stays (Unit 2, Low Risk, and other warehouse names)",
        ],
    )
    picture(doc, d["filters"])
    caption(doc, "Picture 5. Filter order on Global Stocks.")
    add_p(doc, "What “done” looks like. You can narrow the list by product tree, buyer label, and location, without losing search, item, warning, or hide-zero.")
    evidence(doc)

    heading(doc, "5. Req 3 — Column chooser")
    add_p(doc, "The table can show many columns. You choose which ones are visible.")
    picture(doc, d["chooser"])
    caption(doc, "Picture 6. Right-click the header. Tick or untick a column.")
    add_p(doc, "What “done” looks like. Right-click works. Hidden columns come back when you tick them again. The page scrolls sideways for extra columns.")
    evidence(doc)

    heading(doc, "6. Req 3.1 — Add qty when a split column is hidden")
    add_p(doc, "Product code is the identity. Product name stays with it.")
    picture(doc, d["peas"])
    caption(doc, "Picture 7. Peas, two traces, same location → one row.")
    picture(doc, d["salt"])
    caption(doc, "Picture 8. Salt, four suppliers, 5 kg each → 20 kg.")
    picture(doc, d["noloc"])
    caption(doc, "Picture 9. Unit 2 and Low Risk never become one row.")
    add_p(doc, "What “done” looks like. Hide Trace or Supplier and qty adds inside the same location. Two warehouses stay two rows.")
    evidence(doc)

    heading(doc, "7. Req 4 — Future purchase orders on the report")
    add_p(doc, "On the right of Global Stocks, each product row can show open POs. Soonest delivery first.")
    bullets(
        doc,
        [
            "Each PO slot: remaining qty, delivery date, PO number",
            "PO 1, then PO 2, then PO 3, as needed",
            "Only Ordered and Partial POs. Not draft, not cancelled, not fully received",
        ],
    )
    picture(doc, d["po"])
    caption(doc, "Picture 10. How future POs appear on the row.")
    add_p(
        doc,
        "Goods-in rule. If PO 43123 ordered 90 kg and warehouse received 40 kg, the row still shows that PO with 50 kg left. When the last 50 kg is received, that PO disappears and the next future PO moves into first place.",
    )
    picture(doc, d["slide"])
    caption(doc, "Picture 11. Partial receive keeps the PO. Full receive slides the next PO in.")
    add_p(doc, "What “done” looks like. Open POs show. Remaining qty is correct. Finished POs leave. The next PO takes their place.")
    evidence(doc)

    heading(doc, "8. Req 5 — Three quantity columns")
    bullets(
        doc,
        [
            "Current stock qty — already on the page",
            "Future order stock qty — total remaining qty on open POs for that row",
            "Grand total — current + future",
        ],
    )
    add_p(doc, "Example. Peas: 100 kg in stock. 50 kg still on open POs. Grand total 150 kg.")
    picture(doc, d["totals"])
    caption(doc, "Picture 12. How the three qty columns are built.")
    add_p(doc, "What “done” looks like. The three numbers match stock plus remaining PO qty, and they follow the same location split as Req 3.1.")
    evidence(doc)

    heading(doc, "9. Req 6 — Reminder email at 00:10")
    add_p(doc, "Same idea as the closing stock email: a short branded message and a file attached. Addresses are set in the app.")
    picture(doc, d["email"])
    caption(doc, "Picture 13. Night job, newest PO date first.")
    bullets(
        doc,
        [
            "Include: Ordered (nothing received yet) and Partial (some received, not finished)",
            "Exclude: Draft, Received, Cancelled",
            "Sort: newest PO date first, so recent orders sit at the top",
            "This is the full open list each night, not a “only new since yesterday” list",
        ],
    )
    add_p(doc, "What “done” looks like. At 00:10 the named inboxes get the email. The list matches open POs in the system, newest first.")
    evidence(doc)

    heading(doc, "10. Req 7 — Planning Excel, consolidated raw material")
    add_p(doc, "Today the plan Excel lists each requirement on its own row. Peas at one level and peas at another level are two rows.")
    picture(doc, d["excel"])
    caption(doc, "Picture 14. Optional extra sheet. Old sheets stay.")
    add_p(doc, "What “done” looks like. A popup asks if you want the consolidated sheet. If yes, peas 2 kg + 10 kg become 12 kg. Plan Requirements and Working do not change.")
    evidence(doc)

    heading(doc, "11. What we will not break")
    picture(doc, d["safe"])
    caption(doc, "Picture 15. Existing working parts stay.")
    bullets(
        doc,
        [
            "Purchase unit, format and version on Guide",
            "Product timeline (we only add the new field into it)",
            "Global Stocks live feed and current filters",
            "The nine columns you see today remain the default view",
            "Planning Excel sheets Plan Requirements and Working",
            "Closing stock email and its own address list",
        ],
    )

    heading(doc, "12. Sign-off")
    add_p(doc, "This pack describes the work. It is not built yet. Sign when the understanding is accepted. After build, fill the evidence tables with screen recordings.")
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pairs = [
        ("Prepared for management", "Date: 17 Sep 2026"),
        ("Approved by (name / signature)", "Date: ________"),
        ("Ready for developers to build", "Yes / No: ________"),
        ("Evidence pack complete (after build)", "Yes / No: ________"),
    ]
    for i, (a, b) in enumerate(pairs):
        c0, c1 = table.rows[i].cells
        c0.paragraphs[0].clear()
        r0 = c0.paragraphs[0].add_run(a)
        set_run_font(r0, size=11, bold=True, color=NAVY)
        c1.paragraphs[0].clear()
        r1 = c1.paragraphs[0].add_run(b)
        set_run_font(r1, size=11, color=NAVY)
        shade(c0, "FFF7ED")
        set_cell_border(c0)
        set_cell_border(c1)

    add_p(doc, "", space_after=12)
    add_p(
        doc,
        "Editable diagram files (open in excalidraw.com): docs/management/excalidraw/",
        size=9,
        color=SLATE,
    )
    add_p(
        doc,
        "Pictures used in this Word file: docs/management/diagrams/",
        size=9,
        color=SLATE,
    )

    doc.save(DOCX)


def main() -> None:
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    EXCAL_DIR.mkdir(parents=True, exist_ok=True)
    diagrams = build_diagrams()
    build_doc(diagrams)
    print(DOCX)
    print(f"diagrams={len(list(PNG_DIR.glob('*.png')))}")
    print(f"excalidraw={len(list(EXCAL_DIR.glob('*.excalidraw')))}")


if __name__ == "__main__":
    main()
